# -*- coding: utf-8 -*-
"""Scavi con unità numerate per classe (A = strati, B = murature, come a Ba'ja), muri disegnati pietra per
pietra con un file per unità, layer in sistemi di riferimento diversi, reperti e target fotogrammetrici nei
CSV della stazione totale, copie e risultati di programmi nella cartella. Dati sintetici: nessun dato reale."""
import os

import geopandas as gpd
import openpyxl
import pandas as pd
import pytest
from shapely.geometry import box

from stratigrafia3d import importa, Scavo
from stratigrafia3d import sigle as sg
from stratigrafia3d import modifiche as md

CRS = "EPSG:32632"            # zona sbagliata di proposito: lo scavo finto è in Giordania (zona 36N)
X0, Y0 = 3085080.0, 3681100.0


def _pietre(x, y, n=6, lato=0.35, passo=0.45, verticale=False):
    """Un muro di ``n`` pietre quadrate in fila."""
    return [box(x + (0 if verticale else i * passo), y + (i * passo if verticale else 0),
                x + (0 if verticale else i * passo) + lato, y + (i * passo if verticale else 0) + lato)
            for i in range(n)]


def _shp(cartella, nome, geoms, crs=CRS, campo="id", valori=None):
    g = gpd.GeoDataFrame({campo: valori or list(range(1, len(geoms) + 1))}, geometry=geoms, crs=CRS)
    if crs != CRS:
        g = g.to_crs(crs)
    g.to_file(os.path.join(cartella, nome + ".shp"))


def _csv(cartella, nome, righe, sep=";"):
    with open(os.path.join(cartella, nome), "w", encoding="utf-8") as f:
        f.write(sep.join(["NAME", "Y", "X", "Z"]) + "\n")
        for r in righe:
            f.write(sep.join(str(v) for v in r) + "\n")


@pytest.fixture(scope="module")
def cartella(tmp_path_factory):
    d = tmp_path_factory.mktemp("classi") / "Scavo per classi"
    for sub in ("SHP", "CSV", "MATRIX", "REPORT", "MATRIX/run_20261008_164419"):
        os.makedirs(d / sub, exist_ok=True)
    # schede e rapporti: A = strati, B = muri; GR = stanze (non sono unità)
    righe = [
        ("A1", "Rubble layer outside B1", "GR2", "", "", "A2", "", "B1", ""),
        ("A2", "Floor fill of room GR2", "GR2", "", "GR2", "A3", "A1", "B1", ""),
        ("A3", "Collapse inside GR2", "GR2", "", "", "", "A2", "", ""),
        ("A4", "Sediment near wall B2", "GR3", "", "", "", "", "B2", ""),
        ("B1", "Large containing wall in Area G", "GR2", "", "", "", "", "", "B2"),
        ("B2", "Wall of room GR3", "GR3", "", "", "", "", "", "B1"),
        ("B3", "Wall continuation of B1", "GR2", "", "", "", "", "", ""),
    ]
    cols = ["US_ID", "Description", "Room/Space", "Filled by", "Filling", "Covering", "Covered by", "Abutting",
            "Bonded with"]
    pd.DataFrame(righe, columns=cols).to_excel(d / "MATRIX" / "Strat_rel.xlsx", index=False)
    vecchia = pd.DataFrame(righe[:4], columns=cols)
    vecchia.to_excel(d / "REPORT" / "Strat_rel_bcp.xlsx", index=False)
    os.utime(d / "REPORT" / "Strat_rel_bcp.xlsx", (1.6e9, 1.6e9))
    # muri: un file per unità, il campo «id» numera le pietre
    _shp(str(d / "SHP"), "B1", _pietre(X0, Y0))
    _shp(str(d / "SHP"), "B2", _pietre(X0 + 3.5, Y0 + 0.6, verticale=True), crs="EPSG:32640")   # altro CRS
    _shp(str(d / "SHP"), "B3", _pietre(X0 + 10, Y0 + 10))                       # disegno sbagliato…
    _shp(str(d / "SHP"), "B3giusto", _pietre(X0, Y0 + 3.0), campo="Num_stones")  # …e quello corretto
    _shp(str(d / "SHP"), "B1_vecchio_EPSG4326", _pietre(X0, Y0), crs="EPSG:4326")
    # un deposito con la sua pianta, nominata stanza.unità
    gpd.GeoDataFrame({"id": [1]}, geometry=[box(X0 + 0.2, Y0 + 0.6, X0 + 2.5, Y0 + 2.6)], crs=CRS).to_file(
        str(d / "SHP" / "GR2.A2.shp"))
    # punti della stazione totale: reperti, target, una copia
    _csv(str(d / "CSV"), "BJ051026_LTA.csv", [(f"LTA{i}", Y0 + 1 + i * 0.2, X0 + 1, 1174.2) for i in range(1, 4)])
    _csv(str(d / "CSV"), "BJ051026_BON.csv", [("BON7", Y0 + 1.5, X0 + 1.5, 1174.1)])
    _csv(str(d / "CSV"), "BJ051026_mire_GR2.A2.csv", [("M.az", Y0, X0, 1175.0), ("M.ro", Y0 + 2, X0 + 2, 1175.1)])
    _csv(str(d / "CSV"), "BJ051026_LTA - Copia.csv", [("LTA1", Y0 + 1.2, X0 + 1, 1174.2)])
    _csv(str(d / "CSV"), "2110_MURI.csv", [("B1a", Y0 + 0.2, X0 + 0.2, 1175.4), ("B1b", Y0 + 0.2, X0 + 2.0, 1175.5),
                                            ("B2a", Y0 + 1.0, X0 + 3.7, 1175.3)], sep=",")
    # inventario dei reperti con la loro US
    pd.DataFrame({"NAME": ["LTA1", "LTA2", "BON7"], "Y": [Y0 + 1.2, Y0 + 1.4, Y0 + 1.5], "X": [X0 + 1] * 2 + [X0 + 1.5],
                  "Z": [1174.2, 1174.2, 1174.1], "LOCUS": ["GR2"] * 3, "US": ["A2", "A2", "A3"]}).to_excel(
        d / "REPORT" / "BJ_COORDINATED_LTA.xlsx", index=False)
    # risultati di un programma (matrix già elaborato): non vanno usati al posto delle schede
    pd.DataFrame({"From": ["A1"], "To": ["A2"], "Relation": ["covers"]}).to_csv(
        d / "MATRIX" / "run_20261008_164419" / "chronological_edges.csv", index=False)
    return str(d)


@pytest.fixture(scope="module")
def abb(cartella):
    return importa.proponi([cartella])


@pytest.fixture(scope="module")
def scavo(abb):
    return importa.applica(abb)


# ------------------------------------------------------------------ codifica delle sigle
def test_codifica():
    c = sg.Codifica.da_valori(["A1", "A2", "B1", "B12", "US.A5"])
    assert c.classi == {"A": 1000, "B": 2000}
    assert c.codice("A1") == 1001 and c.codice("B1") == 2001 and c.codice("US.A5") == 1005
    assert c.codice("USA10") == 1010 and c.codice(12) == 12 and c.codice("GR2") is None
    assert c.sigla(2012) == "B12" and c.sigla(5) is None
    assert c.nel_testo("B3, B4; GR2, A 12") == ([2003, 2004, 1012], ["GR2"])
    # senza numeri ripetuti tra classi nulla cambia
    assert sg.Codifica.da_valori(["US 1001", "US 1002", 1003]) is None
    assert sg.Codifica.da_valori(["A1", "A2"]) is None


def test_sigle_senza_ritorni_esponenziali():
    import time
    t = time.time()
    for v in ["x" * 5000, "Large containing wall in Area G " * 50, "1" * 30 + "a" * 30]:
        sg.parti(v)
    assert time.time() - t < 0.5


def test_unita_nel_nome():
    c = sg.Codifica({"A": 1000, "B": 2000}, 1000)
    noti = {1002, 1066, 2001, 2003, 2012}
    assert sg.unita_nel_nome("B12", c, noti) == [2012]
    assert sg.unita_nel_nome("B3giusto", c, noti) == [2003]
    assert sg.unita_nel_nome("GR6.A66", c, noti) == [1066]
    assert sg.unita_nel_nome("G2.POL2", c, noti) == []


# ------------------------------------------------------------------ inventario
def test_inventario_scarta_copie_derivati_e_versioni(abb):
    voci = {v["relativo"]: v for v in abb.inventario["voci"]}
    assert voci["CSV/BJ051026_LTA - Copia.csv"]["usa"] is False
    assert voci["MATRIX/run_20261008_164419/chronological_edges.csv"]["usa"] is False
    assert voci["REPORT/Strat_rel_bcp.xlsx"]["usa"] is False
    assert voci["MATRIX/Strat_rel.xlsx"]["usa"] is not False
    assert abb.tabella.endswith(os.path.join("MATRIX", "Strat_rel.xlsx"))


# ------------------------------------------------------------------ proposta
def test_proposta(abb):
    assert abb.sigle and set(abb.sigle["classi"]) == {"A", "B"}
    r = {x.layer: x for x in abb.layers}
    assert r["B1"].ruolo == "usm" and r["B1"].unita_nome == 2001
    assert r["B2"].ruolo == "usm" and r["B2"].unita_nome == 2002
    assert r["B3"].ruolo == "ignora" and "B3giusto" in r["B3"].motivo
    assert r["B3giusto"].ruolo == "usm" and r["B3giusto"].unita_nome == 2003
    assert r["GR2.A2"].ruolo == "us" and r["GR2.A2"].unita_nome == 1002
    assert abb.classi_murarie == ["B"]
    assert r["BJ051026_LTA.csv"].ruolo == "reperti" and r["BJ051026_LTA.csv"].classe_reperto == "LTA"
    # i target della fotogrammetria non diventano quote (l'inventario non li propone nemmeno)
    assert r.get("BJ051026_mire_GR2.A2.csv") is None or r["BJ051026_mire_GR2.A2.csv"].ruolo == "ignora"
    voci = {v["relativo"]: v for v in abb.inventario["voci"]}
    assert voci["CSV/BJ051026_mire_GR2.A2.csv"]["usa"] is False
    # crs di lavoro: il più usato tra quelli in metri, non quello geografico
    assert abb.crs == CRS
    assert any("zona 36N" in n for n in abb.note)
    assert any("Covering" in n for n in abb.note)          # le colonne in -ing sono rapporti


# ------------------------------------------------------------------ progetto
def test_progetto(scavo):
    us, usm = scavo.schede_us(), scavo.schede_usm()
    assert set(us) == {1001, 1002, 1003, 1004}
    assert set(usm) == {2001, 2002, 2003}
    et = scavo.etichette()
    assert et[1001] == "A1" and et[2001] == "B1"
    rap = scavo.tabelle["Rapporti"]
    coppie = {(a, t, b) for a, t, b in rap.itertuples(index=False)}
    assert (1001, "copre", 1002) in coppie                   # A1 copre A2, non «1 copre 2»
    assert (1001, "si appoggia a", 2001) in coppie           # Abutting
    assert not any(b < 1000 for _, _, b in coppie)           # «GR2» non diventa l'unità 2
    problemi = scavo.verifica()
    assert not [p for p in problemi if p.livello == "errore"], [str(p) for p in problemi]
    # i messaggi mostrano le sigle
    assert all(not p.unita or p.etichette for p in problemi)
    # muri nel CRS di lavoro, vicini tra loro anche se B2 era in un altro sistema
    g = scavo.layers["usm_poligoni"]
    assert g.crs.to_string() == CRS
    x0, y0, x1, y1 = g.total_bounds
    assert x1 - x0 < 20 and y1 - y0 < 20


def test_reperti_con_us_dall_inventario(scavo):
    rs = scavo.tabelle["Reperti_speciali"]
    assert {"LTA1", "LTA2", "LTA3", "BON7"} <= set(rs["RS"])
    per = dict(zip(rs["RS"], rs["US"]))
    assert per["LTA1"] == 1002 and per["BON7"] == 1003
    assert set(rs.loc[rs.RS == "LTA1", "Materiale"]) == {"LTA"}


def test_ricostruzione_e_visualizzatore(scavo, tmp_path):
    from stratigrafia3d.ricostruzione import ricostruisci
    from stratigrafia3d.esporta import pagina_visualizzatore
    ricostruisci(scavo)
    assert 2001 in scavo.modello.unita and 1002 in scavo.modello.unita
    html = pagina_visualizzatore(scavo)
    assert '"Sigla":"B1"' in html and "const ETI=" in html


def test_elaborati_con_sigle(scavo, tmp_path):
    from stratigrafia3d.ricostruzione import ricostruisci
    from stratigrafia3d import elaborati
    if scavo.modello is None:
        ricostruisci(scavo)
    t = elaborati.tabella_volumi(scavo)
    assert "Sigla" in t.columns and "B1" in set(t["Sigla"])
    p = tmp_path / "pianta.svg"
    elaborati.pianta_svg(scavo, str(p))
    assert ">B1<" in p.read_text(encoding="utf-8") or "B1" in p.read_text(encoding="utf-8")


def test_riscrittura_con_sigle(cartella, tmp_path):
    """Un rapporto aggiunto nell'app si scrive con le sigle («B3»), lasciando le stanze («GR2») nella cella."""
    import shutil
    d = tmp_path / "copia"
    shutil.copytree(cartella, d)
    s = importa.applica(importa.proponi([str(d)]))
    md.applica_modifica(s, 1002, aggiungi=[("si appoggia a", 2003)])
    md.riscrivi(s, copie=str(tmp_path / "copie"))
    wb = openpyxl.load_workbook(d / "MATRIX" / "Strat_rel.xlsx")
    ws = wb.active
    intest = [c.value for c in ws[1]]
    riga = next(r for r in ws.iter_rows(min_row=2, values_only=True) if r[0] == "A2")
    abut = riga[intest.index("Abutting")]
    assert "B3" in str(abut) and "2003" not in str(abut)
    riempie = riga[intest.index("Filling")]
    assert riempie in (None, "", "GR2") or "GR2" in str(riempie)


def test_classi_murarie_a_mano(cartella):
    abb = importa.proponi([cartella])
    importa.imposta_classi_murarie(abb, [])
    r = {x.layer: x for x in abb.layers}
    assert r["B1"].ruolo == "us" and abb.classi_murarie == []
    importa.imposta_classi_murarie(abb, ["b"])
    assert r["B1"].ruolo == "usm" and abb.classi_murarie == ["B"]


def test_profilo_conserva_le_sigle(abb, tmp_path):
    p = tmp_path / "profilo.json"
    abb.salva_profilo(str(p))
    b = importa.Abbinamento.carica_profilo(str(p))
    assert b.sigle == abb.sigle and b.classi_murarie == ["B"]
    assert {x.layer: x.unita_nome for x in b.layers}["B1"] == 2001
