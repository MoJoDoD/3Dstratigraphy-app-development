"""Modifiche nell'app, riscrittura nei file d'origine (Excel, CSV, SQLite) e aggiornamento dai file."""
import ast
import os
import shutil
import sqlite3

import geopandas as gpd
import openpyxl
import pandas as pd
import pyogrio
import pytest
from shapely.geometry import Point, box

from stratigrafia3d import importa, ricostruisci, Scavo
from stratigrafia3d import modifiche as md


@pytest.fixture
def demo_copia(demo, tmp_path):
    g = shutil.copy(demo["gpkg"], tmp_path / "scavo.gpkg")
    x = shutil.copy(demo["xlsx"], tmp_path / "scavo.xlsx")
    s = importa.applica(importa.proponi([str(g), str(x)]))
    return s, str(x), tmp_path


def _cella(xlsx, foglio, col_id, u, col):
    df = pd.read_excel(xlsx, sheet_name=foglio)
    return df.loc[df[col_id] == u, col].iloc[0]


def test_modifica_e_riscrittura_excel(demo_copia):
    s, xlsx, d = demo_copia
    stile = openpyxl.load_workbook(xlsx)["US"]["A1"].font.b
    r = md.applica_modifica(s, 1005, {"Descrizione": "Tegole in crollo, riviste", "Spessore medio stimato (m)": "0,25"},
                            aggiungi=[("copre", 1035)], togli=[("coperto da", 1000)])
    assert r["voci"] == 4 and 1005 in r["ricostruire"]
    assert r["scheda"]["Spessore medio stimato (m)"] == 0.25
    assert [1000, "copre", 1005] not in r["rapporti"] and [1005, "copre", 1035] in r["rapporti"]
    assert len(md.in_attesa(s)) == 4
    # salvato e riaperto, il registro resta
    p = str(d / "p.scavo")
    s.salva(p)
    s = Scavo.apri(p)
    assert len(md.in_attesa(s)) == 4
    esito = md.riscrivi(s, copie=str(d / "copie"))
    assert esito["scritte"] == 4 and not esito["saltate"], esito["saltate"]
    assert len(esito["copie"]) == 1 and os.path.exists(esito["copie"][0])
    assert _cella(xlsx, "US", "US", 1005, "Descrizione") == "Tegole in crollo, riviste"
    assert _cella(xlsx, "US", "US", 1005, "Spessore medio stimato (m)") == 0.25
    rap = pd.read_excel(xlsx, sheet_name="Rapporti")
    righe = set(map(tuple, rap.iloc[:, :3].values.tolist()))
    assert (1005, "copre", 1035) in righe and (1000, "copre", 1005) not in righe
    assert openpyxl.load_workbook(xlsx)["US"]["A1"].font.b == stile        # formattazione conservata
    assert not md.in_attesa(s) and not md.sorgenti_cambiate(s)
    # rileggendo il file con la stessa ricetta non cambia nulla
    nuovo, ricostruire, note = md.ricarica(s)
    assert ricostruire == [] and note == ["Nessuna unità cambiata"]


def test_rapporto_che_crea_un_ciclo(demo_copia):
    s = demo_copia[0]
    with pytest.raises(md.ErroreModifica, match="ciclo"):
        md.applica_modifica(s, 1000, aggiungi=[("coperto da", 1005)])
    assert not md.in_attesa(s)
    with pytest.raises(md.ErroreModifica, match="non esiste"):
        md.applica_modifica(s, 1000, aggiungi=[("copre", 99999)])


def test_aggiornamento_dai_file(demo_copia):
    s, xlsx, d = demo_copia
    ricostruisci(s)
    wb = openpyxl.load_workbook(xlsx)
    ws = wb["US"]
    col = [c.value for c in ws[1]]
    for row in ws.iter_rows(min_row=2):
        if row[col.index("US")].value == 1011:
            row[col.index("Spessore medio stimato (m)")].value = 0.33
    wb.save(xlsx)
    assert [x["stato"] for x in md.sorgenti_cambiate(s)] == ["cambiato"]
    md.applica_modifica(s, 1005, {"Descrizione": "x"})
    with pytest.raises(md.ErroreModifica, match="non ancora scritte"):
        md.ricarica(s)
    nuovo, ricostruire, note = md.ricarica(s, scarta_modifiche=True)
    tutte = set(s.modello.unita)
    assert 1011 in ricostruire and len(ricostruire) < len(tutte)
    assert nuovo.schede_us()[1011]["Spessore medio stimato (m)"] == 0.33
    assert nuovo.schede_us()[1005]["Descrizione"] != "x"
    assert nuovo.modello is not None
    ricostruisci(nuovo, unita=ricostruire)
    assert set(nuovo.modello.unita) == set(s.modello.unita)


def test_riscrittura_csv_con_vocabolario_e_colonna_padre(tmp_path):
    g = gpd.GeoDataFrame({"CONTEXT_ID": [10, 20]}, geometry=[box(0, 0, 4, 2), box(6, 0, 7, 1)], crs="EPSG:27700")
    g.to_file(str(tmp_path / "Stansted.shp"))
    csv = tmp_path / "ContextData.csv"
    pd.DataFrame({"Context Number": [10, 11, 20, 21], "Context Type": ["Cut", "Deposit", "Cut", "Deposit"],
                  "Fill of": [None, 10, None, None], "Interpretation": ["Ditch", "Fill – café", "Pit", "Fill"],
                  "Context Depth (m)": [0.8, 0.3, None, 0.2]}).to_csv(csv, index=False, sep=";", encoding="cp1252")
    abb = importa.proponi([str(tmp_path / "Stansted.shp"), str(csv)])
    abb, _ = importa.applica_ricetta(abb, importa.carica_ricetta_pronta("framework_archaeology"))
    abb.solo_con_poligono = False
    s = importa.applica(abb)
    md.applica_modifica(s, 21, {"Tipo": "negativa", "Spessore/profondità max (m)": 0.45}, aggiungi=[("riempie", 20)])
    md.applica_modifica(s, 11, togli=[("riempie", 10)])
    esito = md.riscrivi(s, copie=str(tmp_path / "copie"))
    assert not esito["saltate"], esito["saltate"]
    testo = open(csv, encoding="cp1252").read()
    assert "café" in testo and ";" in testo.splitlines()[0]
    df = pd.read_csv(csv, sep=";", encoding="cp1252").set_index("Context Number")
    assert df.loc[21, "Context Type"] == "Cut" and df.loc[21, "Context Depth (m)"] == 0.45
    assert df.loc[21, "Fill of"] == 20 and pd.isna(df.loc[11, "Fill of"])


def test_riscrittura_pyarchinit(tmp_path):
    db = str(tmp_path / "pyarchinit_db.sqlite")
    us = gpd.GeoDataFrame({"scavo_s": ["Sito"] * 3, "us_s": [1, 2, 3]},
                          geometry=[box(0, 0, 4, 4), box(0.5, 0.5, 2, 2), box(0, 0, 4, 4)], crs="EPSG:3004")
    pyogrio.write_dataframe(us, db, layer="pyunitastratigrafiche", driver="SQLite", dataset_options={"SPATIALITE": "YES"})
    q = gpd.GeoDataFrame({"us_q": [1, 1, 1, 3, 3, 3], "quota_q": [10.0, 10.1, 10.05, 9.6, 9.55, 9.6]},
                         geometry=[Point(1, 1), Point(3, 1), Point(2, 3)] * 2, crs="EPSG:3004")
    pyogrio.write_dataframe(q, db, layer="pyarchinit_quote", driver="SQLite", append=True)
    tab = pd.DataFrame({"sito": ["Sito"] * 3, "area": ["1"] * 3, "us": [1, 2, 3],
                        "d_stratigrafica": ["Strato", "Taglio", "Strato"], "descrizione": ["a", "b", "c"],
                        "rapporti": ["[['Copre', '3', '1', 'Sito']]", "[['Taglia', '3', '1', 'Sito']]",
                                     "[['Coperto da', '1', '1', 'Sito'], ['Tagliato da', '2', '1', 'Sito']]"]})
    pyogrio.write_dataframe(tab, db, layer="us_table", driver="SQLite", append=True)
    abb, _ = importa.applica_ricetta(importa.proponi([db]), importa.carica_ricetta_pronta("pyarchinit"))
    s = importa.applica(abb)
    md.applica_modifica(s, 1, {"Descrizione": "macerie"}, aggiungi=[("copre", 2)])
    esito = md.riscrivi(s, copie=str(tmp_path / "copie"))
    assert esito["scritte"] == 2 and not esito["saltate"], esito["saltate"]
    con = sqlite3.connect(db)
    descr, rap = con.execute("SELECT descrizione, rapporti FROM us_table WHERE us = 1").fetchone()
    assert descr == "macerie"
    assert ast.literal_eval(rap) == [["Copre", "3", "1", "Sito"], ["Copre", "2", "1", "Sito"]]
