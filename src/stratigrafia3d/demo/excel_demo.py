# -*- coding: utf-8 -*-
"""Database di scavo in Excel (openpyxl), con fogli collegati da formule."""
import datetime
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.comments import Comment

F = "Arial"
HFILL = PatternFill("solid", fgColor="3B3024")
HFONT = Font(name=F, bold=True, color="FFFFFF", size=10)
BFONT = Font(name=F, size=9)
TFONT = Font(name=F, bold=True, size=14)
thin = Side(style="thin", color="D9D2C5")


def _sheet(wb, name, headers, rows, widths=None, wrap_cols=(), table=True, note=None):
    ws = wb.create_sheet(name)
    ws.append(headers)
    for r in rows:
        ws.append([None if (isinstance(v, float) and np.isnan(v)) else v for v in r])
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = HFONT; cell.fill = HFILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.font = BFONT
            cell.alignment = Alignment(vertical="top", wrap_text=(cell.column in wrap_cols))
    for i, w in enumerate(widths or [], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "B2"
    if table and ws.max_row > 1:
        ref = f"A1:{get_column_letter(len(headers))}{ws.max_row}"
        t = Table(displayName="T_" + name.replace(" ", "_"), ref=ref)
        t.tableStyleInfo = TableStyleInfo(name="TableStyleLight15", showRowStripes=True)
        ws.add_table(t)
    if note:
        ws.cell(row=1, column=1).comment = Comment(note, "Database")
    return ws


def scrivi_excel(path, D):
    US, USM, FASI = D["US"], D["USM"], D["FASI"]
    fmt = D["fmt_date"]
    fase_d = {f[0]: f for f in FASI}
    wb = Workbook()
    ws = wb.active; ws.title = "Info"
    S = D["SITO"]
    ws["A1"] = "Database di scavo – " + S["nome"]; ws["A1"].font = TFONT
    ws["A2"] = S["note"]; ws["A2"].font = Font(name=F, bold=True, color="B00000", size=10)
    info = [("Area", S["area"]), ("Campagna", S["campagna"]), ("Sistema di riferimento", S["crs"]),
            ("Origine griglia locale", S["origine_locale"]), ("Quote", S["quote"]),
            ("GIS collegato", "01_GIS/podere_roveto_area1000.gpkg (chiave di collegamento: numero di US/USM)"), ("", ""),
            ("Fogli", ""),
            ("US", "Schede delle unità stratigrafiche (positive e negative) con rapporti"),
            ("USM", "Schede delle unità stratigrafiche murarie"),
            ("Fasi", "Periodizzazione"),
            ("Rapporti", "Rapporti stratigrafici diretti (uno per riga) – base del diagramma di Harris"),
            ("Materiali", "Inventario dei materiali per US e classe"),
            ("Reperti_speciali", "Reperti speciali (RS) posizionati in 3D"),
            ("Campioni", "Campionature e risultati delle analisi (fittizi)"),
            ("Quote", "Quote rilevate con stazione totale (stesse del layer GIS 'quote')"),
            ("Sezioni", "Sezioni stratigrafiche rilevate"),
            ("Sintesi_US", "Riepilogo calcolato con formule dagli altri fogli"),
            ("Documentazione", "Registro di foto e disegni"), ("", ""),
            ("Convenzioni", ""),
            ("Datazioni", "Anni numerici: negativi = a.C., positivi = d.C."),
            ("Tipi di quota", "sup = superficie superiore della US; inf = superficie inferiore (dopo l'asportazione); "
                              "taglio = superficie di un'interfaccia negativa; orlo = limite superiore del taglio; "
                              "rasatura / fondazione = quote murarie"),
            ("Stato scavo", "sì = asportata; parziale (saggio) = asportata solo nel Saggio 1; no = lasciata in posto")]
    for i, (k, v) in enumerate(info, start=4):
        ws.cell(row=i, column=1, value=k).font = Font(name=F, bold=True, size=10)
        ws.cell(row=i, column=2, value=v).font = Font(name=F, size=10)
    ws.column_dimensions["A"].width = 26; ws.column_dimensions["B"].width = 110

    # ---- Fasi
    rows = []
    for i, f in enumerate(FASI):
        r = i + 2
        rows.append([f[0], f[1], f[2], f[3], fmt(f[2]), fmt(f[3]), f[4], f[5],
                     f'=COUNTIF(US!$F:$F,A{r})+COUNTIF(USM!$D:$D,A{r})'])
    _sheet(wb, "Fasi", ["Fase", "Periodo", "Da (anno)", "A (anno)", "Da", "A", "Titolo", "Descrizione", "N. US/USM"],
           rows, [7, 18, 10, 10, 11, 11, 28, 90, 11], wrap_cols=(8,))

    # ---- US
    base_date = datetime.date(2026, 6, 8)
    excav_order = [u for u in reversed(D["order"]) if u in US]
    dscavo = {u: base_date + datetime.timedelta(days=int(i * 0.9) + (2 * (i // 5))) for i, u in enumerate(excav_order)}
    resp = ["M.R.", "L.B.", "G.F.", "A.C."]
    rel = D["rel_by_us"]
    J = lambda u, k: ", ".join(str(x) for x in sorted(rel.get(u, {}).get(k, [])))
    rows = []
    for i, u in enumerate(sorted(US)):
        d = US[u]; st = D["us_stats"][u]; f = fase_d[d["fase"]]; r = i + 2
        rows.append([u, d["tipo"], d["cat"], d["defin"], d["amb"], d["fase"], f[1], f[2], f[3],
                     d["munsell"], d["hex"], d["consist"], d["comp"], d["descr"], d["interp"], d["criteri"],
                     d["morf"], d["margini"], D["stato"](u), st["area"],
                     f'=_xlfn.MAXIFS(Quote!$F:$F,Quote!$B:$B,A{r})', f'=_xlfn.MINIFS(Quote!$F:$F,Quote!$B:$B,A{r})',
                     st["sp_medio"], st["sp_max"],
                     J(u, "copre"), J(u, "coperto da"), J(u, "taglia"), J(u, "tagliato da"), J(u, "riempie"),
                     J(u, "riempito da"), J(u, "si appoggia a"), J(u, "gli si appoggia"),
                     "Saggio 1" if u in D["only_sond"] else "Area 1000",
                     dscavo[u], resp[i % 4]])
    hdr = ["US", "Tipo", "Categoria", "Definizione", "Ambiente", "Fase", "Periodo", "Datazione da", "Datazione a",
           "Colore Munsell", "Colore HEX", "Consistenza", "Componenti", "Descrizione", "Interpretazione",
           "Criteri di distinzione", "Morfologia", "Margini", "Stato scavo", "Area documentata (m²)",
           "Quota max (m)", "Quota min (m)", "Spessore medio stimato (m)", "Spessore/profondità max (m)",
           "Copre", "Coperto da", "Taglia", "Tagliato da", "Riempie", "Riempito da", "Si appoggia a",
           "Gli si appoggia", "Settore documentato", "Data scavo", "Responsabile"]
    ws = _sheet(wb, "US", hdr, rows,
                [7, 9, 14, 30, 16, 6, 16, 9, 9, 11, 10, 12, 30, 55, 36, 26, 13, 11, 14, 10, 9, 9, 10, 10,
                 16, 16, 16, 16, 10, 12, 16, 16, 12, 11, 9], wrap_cols=(4, 13, 14, 15, 16, 25, 26, 27, 28, 31, 32))
    for row in ws.iter_rows(min_row=2, min_col=34, max_col=34):
        for c in row:
            c.number_format = "dd/mm/yyyy"
    for row in ws.iter_rows(min_row=2, min_col=21, max_col=24):
        for c in row:
            c.number_format = "0.000" if c.column <= 22 else "0.00"
    ws.cell(row=1, column=23).comment = Comment(
        "Stima dallo spessore osservato in sezione e nelle quote sup/inf. "
        "Usata dal visualizzatore 3D per le US non asportate (base non rilevata).", "Database")
    ws.cell(row=1, column=21).comment = Comment("Formula: massimo delle quote del foglio Quote per questa US.", "Database")

    # ---- USM
    rows = []
    for i, w in enumerate(sorted(USM)):
        d = USM[w]; st = D["wall_stats"][w]; f = fase_d[d["fase"]]; r = i + 2
        rows.append([w, d["cat"], d["defin"], d["fase"], f[1], d["tecnica"], d["materiali"], d["orient"],
                     st["lungh"], st["largh"], st["area"],
                     f'=_xlfn.MAXIFS(Quote!$F:$F,Quote!$B:$B,A{r},Quote!$C:$C,"rasatura")',
                     f'=_xlfn.MINIFS(Quote!$F:$F,Quote!$B:$B,A{r},Quote!$C:$C,"rasatura")',
                     st["q_fond"], st["q_fond_stimata"] if st["q_fond"] is None else st["q_fond"],
                     f'=AVERAGEIFS(Quote!$F:$F,Quote!$B:$B,A{r},Quote!$C:$C,"rasatura")-O{r}',
                     d["descr"], d["interp"],
                     J(w, "si lega a"), J(w, "gli si appoggia"), J(w, "si appoggia a"), J(w, "riempie"),
                     J(w, "tagliato da"), J(w, "coperto da")])
    ws = _sheet(wb, "USM", ["USM", "Categoria", "Definizione", "Fase", "Periodo", "Tecnica costruttiva", "Materiali",
                            "Orientamento", "Lunghezza (m)", "Larghezza (m)", "Area (m²)", "Quota rasatura max",
                            "Quota rasatura min", "Quota fondazione rilevata", "Quota base usata (rilevata o stimata)",
                            "Altezza conservata media (m)", "Descrizione", "Interpretazione", "Si lega a",
                            "Gli si appoggia", "Si appoggia a", "Riempie", "Tagliata da", "Coperta da"],
                rows, [7, 11, 30, 6, 14, 40, 30, 10, 10, 10, 9, 10, 10, 11, 12, 11, 50, 34, 10, 26, 12, 9, 10, 12],
                wrap_cols=(3, 6, 7, 17, 18, 20))
    for row in ws.iter_rows(min_row=2, min_col=12, max_col=16):
        for c in row:
            c.number_format = "0.000" if c.column < 16 else "0.00"
    ws.cell(row=1, column=15).comment = Comment(
        "Dove la fondazione non è stata raggiunta, quota stimata per confronto con i muri visti nel saggio.", "Database")

    # ---- Rapporti
    tipo = lambda u: "USM" if u in USM else "US"
    rows = [[a, t, b, tipo(a), tipo(b)] for a, t, b in D["rel"]]
    _sheet(wb, "Rapporti", ["US", "Rapporto", "US correlata", "Tipo", "Tipo correlata"], rows, [8, 16, 12, 8, 12])

    # ---- Materiali
    m = D["mdf"]
    rows = [[r.id, r.us, r.cassetta, r.classe, r.tipo_forma, r.NR, r.NMI, r.peso_g, r.datazione_da, r.datazione_a]
            for r in m.itertuples()]
    ws = _sheet(wb, "Materiali", ["ID", "US", "Cassetta", "Classe", "Tipo / forma", "NR", "NMI", "Peso (g)",
                                  "Datazione da", "Datazione a"], rows, [9, 7, 9, 26, 46, 7, 7, 10, 10, 10],
                wrap_cols=(5,))
    for row in ws.iter_rows(min_row=2, min_col=8, max_col=8):
        for c in row:
            c.number_format = "#,##0.0"

    # ---- Reperti speciali
    E0, N0 = D["E0"], D["N0"]
    rows = [[r.rs, r.us, r.oggetto, r.materiale, r.descrizione, r.dimensioni, r.datazione_da, r.datazione_a,
             round(r.x + E0, 3), round(r.y + N0, 3), r.z] for r in D["rsdf"].itertuples()]
    ws = _sheet(wb, "Reperti_speciali", ["RS", "US", "Oggetto", "Materiale", "Descrizione", "Dimensioni", "Datazione da",
                                         "Datazione a", "E (m)", "N (m)", "Quota (m)"], rows,
                [7, 7, 16, 12, 60, 14, 10, 10, 13, 13, 9], wrap_cols=(5,))
    for row in ws.iter_rows(min_row=2, min_col=9, max_col=11):
        for c in row:
            c.number_format = "0.000"

    # ---- Campioni
    rows = [[r.campione, r.us, r.tipo, r.analisi, r.risultato, round(r.x + E0, 3), round(r.y + N0, 3), r.z]
            for r in D["cdf"].itertuples()]
    _sheet(wb, "Campioni", ["Campione", "US", "Tipo", "Analisi", "Risultato (fittizio)", "E (m)", "N (m)", "Quota (m)"],
           rows, [10, 7, 18, 30, 60, 13, 13, 9], wrap_cols=(5,))

    # ---- Quote
    q = D["qdf"]
    rows = [[r.id, r.us, r.tipo, round(r.x + E0, 3), round(r.y + N0, 3), r.z] for r in q.itertuples()]
    ws = _sheet(wb, "Quote", ["ID", "US", "Tipo", "E (m)", "N (m)", "Quota (m)"], rows, [8, 7, 11, 13, 13, 10])
    for row in ws.iter_rows(min_row=2, min_col=4, max_col=6):
        for c in row:
            c.number_format = "0.000"

    # ---- Sezioni
    prof = D["prof"]
    rows = []
    for s, (p0, p1, ori) in D["SEZIONI"].items():
        uu = sorted(set(int(u) for u in prof[prof.sezione == s].us))
        rows.append([s, ori, round(np.hypot(p1[0] - p0[0], p1[1] - p0[1]), 2),
                     f"{p0[0] + E0:.2f} / {p0[1] + N0:.2f}", f"{p1[0] + E0:.2f} / {p1[1] + N0:.2f}",
                     ", ".join(map(str, uu)), f"Tav. {list(D['SEZIONI']).index(s) + 3} – scala 1:20"])
    _sheet(wb, "Sezioni", ["Sezione", "Orientamento", "Lunghezza (m)", "Inizio E / N", "Fine E / N",
                           "US / USM attraversate", "Disegno"], rows, [9, 12, 12, 22, 22, 70, 20], wrap_cols=(6,))

    # ---- Sintesi (formule)
    rows = []
    allu = sorted(US) + sorted(USM)
    for i, u in enumerate(allu):
        r = i + 2
        cat = US[u]["cat"] if u in US else USM[u]["cat"]
        fa = US[u]["fase"] if u in US else USM[u]["fase"]
        rows.append([u, "US" if u in US else "USM", cat, fa,
                     f"=SUMIFS(Materiali!$F:$F,Materiali!$B:$B,A{r})",
                     f"=SUMIFS(Materiali!$G:$G,Materiali!$B:$B,A{r})",
                     f"=SUMIFS(Materiali!$H:$H,Materiali!$B:$B,A{r})",
                     f"=COUNTIF(Materiali!$B:$B,A{r})",
                     f"=COUNTIF(Reperti_speciali!$B:$B,A{r})",
                     f"=COUNTIF(Campioni!$B:$B,A{r})",
                     f"=COUNTIF(Quote!$B:$B,A{r})",
                     f"=COUNTIF(Rapporti!$A:$A,A{r})+COUNTIF(Rapporti!$C:$C,A{r})"])
    n = len(rows) + 1
    rows.append(["TOTALE", "", "", "", f"=SUM(E2:E{n})", f"=SUM(F2:F{n})", f"=SUM(G2:G{n})", f"=SUM(H2:H{n})",
                 f"=SUM(I2:I{n})", f"=SUM(J2:J{n})", f"=SUM(K2:K{n})", ""])
    ws = _sheet(wb, "Sintesi_US", ["US/USM", "Tipo", "Categoria", "Fase", "NR totale", "NMI totale", "Peso totale (g)",
                                   "Righe di inventario", "Reperti speciali", "Campioni", "Quote", "Rapporti"],
                rows, [9, 7, 16, 6, 10, 10, 13, 11, 10, 9, 8, 9], table=False)
    for c in ws[n + 1]:
        c.font = Font(name=F, bold=True, size=9)
    for row in ws.iter_rows(min_row=2, min_col=7, max_col=7):
        for c in row:
            c.number_format = "#,##0"
    ws.auto_filter.ref = f"A1:L{n}"

    # ---- Documentazione
    rows = []
    k = 1
    for u in sorted(US):
        for v in (["pianta", "dettaglio"] if US[u]["cat"] != "taglio" else ["svuotato"]):
            rows.append([f"F{k:04d}", "Foto", u, f"US {u} – {US[u]['defin'].lower()}, {v}, da {['N', 'S', 'E', 'O'][k % 4]}",
                         dscavo[u].strftime("%d/%m/%Y")])
            k += 1
    for u in sorted(USM):
        rows.append([f"F{k:04d}", "Foto", u, f"USM {u} – {USM[u]['defin'].lower()}, prospetto", "10/07/2026"]); k += 1
    for j, s in enumerate(D["SEZIONI"]):
        rows.append([f"D{j + 3:03d}", "Disegno", "", f"Sezione {s}, scala 1:20", "15/07/2026"])
    rows.append(["D001", "Disegno", "", "Pianta composita Area 1000, scala 1:20", "17/07/2026"])
    rows.append(["D002", "Disegno", "", "Pianta del Saggio 1, scala 1:20", "16/07/2026"])
    rows.append(["R001", "Rilievo", "", "Rilievo con stazione totale di quote, limiti e reperti speciali", "17/07/2026"])
    _sheet(wb, "Documentazione", ["ID", "Tipo", "US/USM", "Soggetto", "Data"], rows, [8, 9, 8, 70, 11], wrap_cols=(4,))

    wb.save(path)
