"""Modifiche alle schede fatte nell'app, riscrittura nei file d'origine e aggiornamento dai file.

Il progetto tiene un registro delle modifiche (``Scavo.modifiche``). Ogni voce dice unità, campo, valore
di prima e di dopo, e se è già stata scritta nei file d'origine. La riscrittura passa per la ricetta
al contrario: la ricetta sa da quale file, foglio e colonna viene ogni campo, quale vocabolario lo ha
tradotto e in quale unità di misura era scritto.

Formati riscrivibili: Excel (.xlsx, con openpyxl: formattazione e formule restano), CSV (stessa
codifica e separatore), tabelle di GeoPackage e SQLite/SpatiaLite. Prima di scrivere, una copia di
ogni file va in ``~/.stratigrafia3d/copie``.
"""
import datetime as _dt
import hashlib
import json
import math
import os
import re
import shutil

import networkx as nx
import numpy as np
import pandas as pd

from . import schema as sc
from . import stratigrafia as st

NUMERICI = {sc.C_SPESSORE, sc.C_PROFONDITA, sc.C_FASE, "Datazione da", "Datazione a", "Quota min (m)",
            "Quota max (m)", "Lunghezza (m)", "Larghezza (m)", "Area documentata (m²)", sc.C_BASE_USM}
# campi che cambiano la forma 3D: dopo una modifica le unità vanno ricostruite
GEOMETRICI = {sc.C_TIPO, sc.C_SPESSORE, sc.C_PROFONDITA, sc.C_FASE, sc.C_BASE_USM}
TUTTI_RAPPORTI = sorted(set(sc.INVERSI) | set(sc.INVERSI.values()))
COPIE = os.path.join(os.path.expanduser("~"), ".stratigrafia3d", "copie")


class ErroreModifica(ValueError):
    pass


def _ora():
    return _dt.datetime.now().isoformat(timespec="seconds")


def _vuoto(v):
    return v is None or (isinstance(v, float) and math.isnan(v)) or (isinstance(v, str) and not v.strip())


def _json(v):
    if _vuoto(v):
        return None
    if isinstance(v, (np.integer,)):
        return int(v)
    if isinstance(v, (np.floating,)):
        return float(v)
    if isinstance(v, (int, float, str, bool)):
        return v
    return str(v)


def _numero(v):
    if _vuoto(v):
        return None
    try:
        return float(str(v).replace(",", ".").strip())
    except ValueError:
        raise ErroreModifica(f"«{v}» non è un numero")


def _diretto(a, t, b):
    """(a, t, b) nella forma diretta: «coperto da» diventa «copre» con i termini scambiati."""
    t = str(t).strip().lower()
    if t in sc.DA_INVERSO:
        return b, sc.DA_INVERSO[t], a
    return a, t, b


def _tabella_di(scavo, u):
    if u in scavo.schede_us():
        return sc.S_US, sc.C_US
    if u in scavo.schede_usm():
        return sc.S_USM, sc.C_USM
    raise ErroreModifica(f"L'unità {u} non ha una scheda nel progetto")


# ---------------------------------------------------------------------------------------------- modifica
def applica_modifica(scavo, unita, campi=None, aggiungi=(), togli=()):
    """Cambia i campi della scheda di ``unita`` e i suoi rapporti.

    ``aggiungi`` e ``togli`` sono coppie (rapporto, altra unità) viste da ``unita``: («copre», 1002).
    Ritorna il dizionario {"ricostruire": [unità], "scheda": {...}, "rapporti": [[a, t, b], ...]}."""
    u = int(unita)
    nome, idc = _tabella_di(scavo, u)
    df = scavo.tabelle[nome]
    riga = df.index[df[idc].map(lambda v: not _vuoto(v) and int(v) == u)]
    if not len(riga):
        raise ErroreModifica(f"Scheda dell'unità {u} non trovata")
    i = riga[0]
    voci, ricostruire = [], set()
    for campo, valore in (campi or {}).items():
        if campo in (idc, sc.C_US, sc.C_USM) or campo.startswith("_"):
            continue
        if campo in NUMERICI:
            valore = _numero(valore)
            if campo == sc.C_FASE and valore is not None:
                valore = int(valore) if float(valore).is_integer() else valore
        elif campo == sc.C_TIPO:
            valore = str(valore or "").strip().lower()
            if valore not in ("positiva", "negativa"):
                raise ErroreModifica("Il tipo deve essere «positiva» o «negativa»")
        else:
            valore = None if _vuoto(valore) else str(valore)
        prima = _json(df.at[i, campo]) if campo in df.columns else None
        if prima == valore or (prima is not None and valore is not None and str(prima) == str(valore)):
            continue
        if campo not in df.columns:
            df[campo] = None
        if df[campo].dtype != object and not (valore is None or isinstance(valore, (int, float))):
            df[campo] = df[campo].astype(object)
        df.at[i, campo] = valore
        voci.append(dict(tipo="campo", unita=u, campo=campo, prima=prima, dopo=valore))
        if campo in GEOMETRICI:
            ricostruire.add(u)

    # rapporti
    rap = scavo.tabelle.get(sc.S_RAPPORTI)
    if rap is None:
        rap = pd.DataFrame(columns=[sc.C_US, "Rapporto", "US correlata"])
    cols = list(rap.columns[:3])
    tutte = set(scavo.schede_us()) | set(scavo.schede_usm())
    norm = [_diretto(int(a), t, int(b)) if not (_vuoto(a) or _vuoto(b)) else None
            for a, t, b in zip(rap[cols[0]], rap[cols[1]], rap[cols[2]])]
    for t, b in togli or ():
        chiave = _diretto(u, t, int(b))
        cont = sc.RAPPORTI_CONTEMPORANEI
        via = [k for k, n in enumerate(norm) if n is not None and (n == chiave or (
            chiave[1] in cont and n[1] in cont and {n[0], n[2]} == {chiave[0], chiave[2]}))]
        if not via:
            raise ErroreModifica(f"Rapporto «{u} {t} {b}» non trovato")
        rap = rap.drop(index=rap.index[via])
        norm = [n for k, n in enumerate(norm) if k not in via]
        voci.append(dict(tipo="rapporto-", unita=u, rapporto=str(t).lower(), altra=int(b)))
        ricostruire |= {u, int(b)}
    for t, b in aggiungi or ():
        t = str(t).strip().lower()
        b = int(b)
        if t not in TUTTI_RAPPORTI:
            raise ErroreModifica(f"Rapporto «{t}» sconosciuto")
        if b == u:
            raise ErroreModifica("Un'unità non può avere un rapporto con se stessa")
        if b not in tutte:
            raise ErroreModifica(f"L'unità {b} non esiste nel progetto")
        if _diretto(u, t, b) in norm:
            continue
        rap = pd.concat([rap, pd.DataFrame([[u, t, b]], columns=cols)], ignore_index=True)
        norm.append(_diretto(u, t, b))
        voci.append(dict(tipo="rapporto+", unita=u, rapporto=t, altra=b))
        ricostruire |= {u, b}
    if any(v["tipo"] != "campo" for v in voci):
        prova = st.da_tabella(rap, list(tutte))
        if not nx.is_directed_acyclic_graph(prova.grafo):
            ciclo = nx.find_cycle(prova.grafo)
            raise ErroreModifica("Il rapporto crea un ciclo nella sequenza: " +
                                 " → ".join(str(a) for a, _ in ciclo) + f" → {ciclo[0][0]}")
        scavo.tabelle[sc.S_RAPPORTI] = rap.reset_index(drop=True)

    quando = _ora()
    for v in voci:
        v.update(quando=quando, scritto=False)
    scavo.modifiche.extend(voci)
    if voci:
        scavo.registra("modifica", dict(unita=u, voci=len(voci)))
    # anche le unità sopra quelle cambiate si appoggiano alla loro forma
    if ricostruire and scavo.modello is not None:
        G = scavo.rapporti().grafo
        for x in list(ricostruire):
            if x in G:
                ricostruire |= nx.ancestors(G, x)
    return dict(voci=len(voci), ricostruire=sorted(ricostruire), scheda=scheda_json(scavo, u),
                rapporti=rapporti_di(scavo, u))


def scheda_json(scavo, u):
    nome, idc = _tabella_di(scavo, u)
    df = scavo.tabelle[nome]
    r = df[df[idc].map(lambda v: not _vuoto(v) and int(v) == u)].iloc[0]
    return {k: _json(v) for k, v in r.items()}


def rapporti_di(scavo, u):
    rap = scavo.tabelle.get(sc.S_RAPPORTI)
    if rap is None:
        return []
    c = list(rap.columns[:3])
    out = []
    for a, t, b in zip(rap[c[0]], rap[c[1]], rap[c[2]]):
        if _vuoto(a) or _vuoto(b):
            continue
        a, t, b = _diretto(int(a), t, int(b))
        if u in (a, b):
            out.append([a, sc.R_LEGA if t in sc.RAPPORTI_CONTEMPORANEI else t, b])
    return out


def in_attesa(scavo):
    return [m for m in scavo.modifiche if not m.get("scritto")]


# ---------------------------------------------------------------------------------------------- fogli
def _sha256(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def mappa_fogli(abb):
    """Nome del foglio nel progetto -> (file, nome nel file), come in importa._tutte_le_tabelle."""
    from .importa import leggi_tabelle
    out = {}
    for f in [abb.tabella] + list(abb.tabelle_extra or []):
        if not f or not os.path.exists(f):
            continue
        for nome in leggi_tabelle(f):
            k = nome if nome not in out else f"{os.path.splitext(os.path.basename(f))[0]} · {nome}"
            out[k] = (f, nome)
    return out


class Foglio:
    """Un foglio (o una tabella) di un file d'origine, da leggere e riscrivere cella per cella."""

    def __init__(self, path, nome, risorse=None):
        # risorse: cartelle di lavoro e connessioni condivise tra i fogli dello stesso file
        self.path, self.nome = path, nome
        self.ext = os.path.splitext(path)[1].lower()
        self.cambiato = False
        risorse = {} if risorse is None else risorse
        if self.ext in (".xlsx", ".xlsm"):
            import openpyxl
            if path not in risorse:
                risorse[path] = openpyxl.load_workbook(path)
            self.wb = risorse[path]
            self.ws = self.wb[nome]
            righe = list(self.ws.iter_rows(values_only=True))
            self.intestazione = [str(c) if c is not None else "" for c in (righe[0] if righe else [])]
            self.righe = {i + 2: dict(zip(self.intestazione, r)) for i, r in enumerate(righe[1:])
                          if any(v is not None for v in r)}
        elif self.ext == ".csv":
            from .importa import leggi_csv
            leggi_csv(path)                       # riconosce codifica e separatore
            self.sep, self.cod = _formato_csv(path)
            self.df = pd.read_csv(path, sep=self.sep, encoding=self.cod, dtype=str, keep_default_na=False)
            self.intestazione = list(self.df.columns)
            self.righe = {i: dict(r) for i, r in self.df.iterrows()}
        elif self.ext in (".gpkg", ".sqlite", ".db"):
            import sqlite3
            if path not in risorse:
                risorse[path] = sqlite3.connect(path)
            self.con = risorse[path]
            cur = self.con.execute(f'SELECT rowid, * FROM "{nome}"')
            self.intestazione = [d[0] for d in cur.description][1:]
            self.righe = {r[0]: dict(zip(self.intestazione, r[1:])) for r in cur.fetchall()}
        else:
            raise ErroreModifica(f"Formato non riscrivibile: {os.path.basename(path)}")

    def trova(self, colonna, u, filtri=()):
        """Chiavi delle righe con il numero ``u`` nella colonna, dentro i filtri della ricetta."""
        from .importa import _intero, _norm
        out = []
        for k, r in self.righe.items():
            if _intero(r.get(colonna)) != u:
                continue
            ok = True
            for f in filtri or ():
                if f.get("colonna") in r and f.get("valori"):
                    dentro = _norm(r.get(f["colonna"])) in {_norm(v) for v in f["valori"]}
                    ok &= dentro if f.get("modo", "tieni") == "tieni" else not dentro
            if ok:
                out.append(k)
        return out

    def valore(self, k, colonna):
        return self.righe[k].get(colonna)

    def imposta(self, k, colonna, valore):
        if colonna not in self.intestazione:
            raise ErroreModifica(f"Colonna «{colonna}» assente in «{self.nome}»")
        if self.ext in (".xlsx", ".xlsm"):
            cella = self.ws.cell(row=k, column=self.intestazione.index(colonna) + 1)
            if isinstance(cella.value, str) and cella.value.startswith("="):
                raise ErroreModifica(f"la cella di «{colonna}» contiene una formula: correggi i dati da cui dipende")
            cella.value = valore
        elif self.ext == ".csv":
            self.df.at[k, colonna] = "" if valore is None else _testo_csv(valore)
        else:
            self.con.execute(f'UPDATE "{self.nome}" SET "{colonna}" = ? WHERE rowid = ?', (valore, k))
        self.righe[k][colonna] = valore
        self.cambiato = True

    def aggiungi(self, valori):
        riga = [valori.get(c) for c in self.intestazione]
        if self.ext in (".xlsx", ".xlsm"):
            self.ws.append(riga)
            k = self.ws.max_row
        elif self.ext == ".csv":
            k = (max(self.righe) + 1) if self.righe else 0
            self.df.loc[k] = ["" if v is None else _testo_csv(v) for v in riga]
        else:
            nomi = [c for c in self.intestazione if valori.get(c) is not None]
            cur = self.con.execute(f'INSERT INTO "{self.nome}" ({", ".join(chr(34) + c + chr(34) for c in nomi)}) '
                                   f'VALUES ({", ".join("?" * len(nomi))})', [valori[c] for c in nomi])
            k = cur.lastrowid
        self.righe[k] = dict(zip(self.intestazione, riga))
        self.cambiato = True

    def togli(self, chiavi):
        for k in sorted(chiavi, reverse=True):
            if self.ext in (".xlsx", ".xlsm"):
                self.ws.delete_rows(k)
                self.righe = {(j - 1 if j > k else j): r for j, r in self.righe.items() if j != k}
            elif self.ext == ".csv":
                self.df = self.df.drop(index=k)
                self.righe.pop(k)
            else:
                self.con.execute(f'DELETE FROM "{self.nome}" WHERE rowid = ?', (k,))
                self.righe.pop(k)
            self.cambiato = True

    def salva(self):
        """Scrive il file (per Excel e SQLite: tutto il file, con gli altri fogli aperti)."""
        if self.ext in (".xlsx", ".xlsm"):
            self.wb.save(self.path)
        elif self.ext == ".csv":
            self.df.to_csv(self.path, sep=self.sep, encoding=self.cod, index=False)
        else:
            self.con.commit()


def _formato_csv(path):
    import csv
    campione = open(path, "rb").read(65536)
    cod = "utf-8-sig" if campione.startswith(b"\xef\xbb\xbf") else None
    if cod is None:
        try:
            campione.decode("utf-8")
            cod = "utf-8"
        except UnicodeDecodeError:
            cod = "cp1252"
    testo = campione.decode("latin-1")
    try:
        sep = csv.Sniffer().sniff("\n".join(testo.splitlines()[:20]), delimiters=",;\t|").delimiter
    except csv.Error:
        sep = ","
    return sep, cod


def _testo_csv(v):
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


# ---------------------------------------------------------------------------------------------- riscrittura
def riscrivi(scavo, copie=COPIE):
    """Scrive nei file d'origine le modifiche non ancora scritte. Ritorna un resoconto:
    {"scritte": n, "saltate": [(voce, motivo)], "file": [...], "copie": [...]}."""
    from .importa import _intero, _norm, _vocabolario, _rapporti_da_testo, _rapporto
    abb = scavo.abbinamento
    if abb is None:
        raise ErroreModifica("Il progetto non ricorda da quali file viene: riscrittura impossibile")
    attese = in_attesa(scavo)
    if not attese:
        return dict(scritte=0, saltate=[], file=[], copie=[])
    fogli_map = mappa_fogli(abb)
    aperti, risorse = {}, {}

    def foglio(nome):
        if nome not in fogli_map:
            raise ErroreModifica(f"Il foglio «{nome}» non si trova più nei file d'origine")
        if nome not in aperti:
            aperti[nome] = Foglio(*fogli_map[nome], risorse=risorse)
        return aperti[nome]

    filtri_scheda = [f for f in (abb.filtri or []) if f.get("dove", "scheda") == "scheda"]
    dal_layer = {k for r in abb.layers for k in (r.campi_scheda or {})}
    usm_nomi = set(scavo.schede_usm())
    saltate, scritte = [], 0

    def origine_campo(u, campo):
        """(foglio, colonna id, colonna del campo) per il campo di una scheda."""
        if u in usm_nomi and abb.foglio_usm:
            nome, colmap, idc = abb.foglio_usm, abb.colonne_usm or {}, (abb.colonne_usm or {}).get(sc.C_USM)
        else:
            nome, colmap, idc = abb.foglio_us, abb.colonne_us or {}, (abb.colonne_us or {}).get(sc.C_US)
        if not nome or not idc:
            raise ErroreModifica("le schede non vengono da un foglio (sono state create dai poligoni)")
        if campo == sc.C_FASE and abb.fase_composta:
            raise ErroreModifica("la fase viene da periodo + fase: correggila nel file d'origine")
        col = colmap.get(campo)
        if col is None:
            f = foglio(nome)
            nudo = campo[:-len(" (archivio)")] if campo.endswith(" (archivio)") else campo
            if nudo in f.intestazione and nudo not in colmap.values():
                col = nudo
            elif campo in dal_layer:
                raise ErroreModifica("il campo viene dagli attributi dei poligoni, non dalla scheda")
            else:
                raise ErroreModifica(f"il campo «{campo}» non ha una colonna nel file d'origine")
        return nome, idc, col

    def riga_di(f, idc, u):
        k = f.trova(idc, u, filtri_scheda)
        if len(k) != 1:
            raise ErroreModifica(f"{'nessuna riga' if not k else str(len(k)) + ' righe'} per l'unità {u} "
                                 f"in «{f.nome}»")
        return k[0]

    def valore_archivio(f, col, campo, nuovo, nome_tab):
        """Il valore come lo scrive l'archivio: vocabolario e unità di misura al contrario."""
        if nuovo is None:
            return None
        if campo == sc.C_TIPO:
            voc = _vocabolario(abb, "tipo")
            grezzi = [r.get(col) for r in f.righe.values() if not _vuoto(r.get(col))]
            candidati = [g for g in grezzi if voc.get(_norm(g)) == nuovo] or \
                        [g for g in grezzi if _norm(g) == nuovo]
            if candidati:
                return max(set(candidati), key=candidati.count)
            return nuovo
        if campo in NUMERICI and isinstance(nuovo, (int, float)):
            # stesso rapporto tra archivio e progetto delle altre righe (cm, mm...)
            tab = scavo.tabelle.get(nome_tab)
            idc_p = sc.C_USM if nome_tab == sc.S_USM else sc.C_US
            rapporti_ = []
            if tab is not None and campo in tab.columns:
                proj = {int(a): b for a, b in zip(tab[idc_p], tab[campo]) if not _vuoto(a) and not _vuoto(b)}
                idc_f = (abb.colonne_usm if nome_tab == sc.S_USM and abb.foglio_usm else abb.colonne_us).get(
                    sc.C_USM if nome_tab == sc.S_USM and abb.foglio_usm else sc.C_US)
                for r in f.righe.values():
                    uu, g = _intero(r.get(idc_f)), r.get(col)
                    try:
                        g = float(str(g).replace(",", "."))
                    except (TypeError, ValueError):
                        continue
                    p = proj.get(uu)
                    if p not in (None, 0) and g:
                        rapporti_.append(g / float(p))
            k = float(np.median(rapporti_)) if rapporti_ else 1.0
            k = min((1.0, 10.0, 100.0, 1000.0), key=lambda c: abs(math.log(max(k, 1e-9) / c)))
            v = nuovo * k
            return int(v) if float(v).is_integer() else round(v, 6)
        return nuovo

    # ---------------------------------------------------------------- campi
    for m in attese:
        if m["tipo"] != "campo":
            continue
        try:
            nome, idc, col = origine_campo(m["unita"], m["campo"])
            f = foglio(nome)
            k = riga_di(f, idc, m["unita"])
            tab = sc.S_USM if m["unita"] in usm_nomi else sc.S_US
            f.imposta(k, col, valore_archivio(f, col, m["campo"], m["dopo"], tab))
            m.update(scritto=True, file=os.path.basename(f.path), foglio=f.nome, colonna=col)
            scritte += 1
        except ErroreModifica as e:
            saltate.append((m, str(e)))

    # ---------------------------------------------------------------- rapporti
    voc = _vocabolario(abb, "rapporto")
    r = abb.rapporti or {"modo": "nessuno"}
    idc_us = (abb.colonne_us or {}).get(sc.C_US)

    def parola(t, esistenti):
        """Come l'archivio scrive il rapporto t: dal vocabolario o dalle parole già usate; in un
        archivio in inglese, la parola inglese."""
        from .importa import RAPPORTI_INGLESE
        esistenti = [g for g in esistenti if not _vuoto(g)]
        for g in esistenti:
            if voc.get(_norm(g)) == t or _rapporto(g) == t:
                return g
        for k_, v in voc.items():
            if v == t:
                return k_
        if any(_norm(g) in RAPPORTI_INGLESE for g in esistenti):
            inglesi = [k_ for k_, v in RAPPORTI_INGLESE.items() if v == t]
            if inglesi:
                w = inglesi[0]
                return w.capitalize() if all(str(g)[:1].isupper() for g in esistenti) else w
        return t

    def lista_numeri(testo):
        if _vuoto(testo):
            return []
        # «1002, 1003», «1002; 1003», ma anche «10.0» scritto da un foglio di calcolo
        return [int(float(x)) for x in re.findall(r"\d+(?:\.\d+)?", str(testo)) if float(x).is_integer()]

    def scrivi_lista(vecchio, numeri):
        sep = ";" if (isinstance(vecchio, str) and ";" in vecchio and "," not in vecchio) else ","
        return (sep + " ").join(str(n) for n in numeri) if numeri else None

    for m in attese:
        if m["tipo"] not in ("rapporto+", "rapporto-"):
            continue
        a, t, b = m["unita"], m["rapporto"], m["altra"]
        aggiungi = m["tipo"] == "rapporto+"
        try:
            fatto = False
            # colonna «padre» (Fill of): un solo valore per riga
            for e in abb.rapporti_extra or []:
                rel = e.get("rapporto", sc.R_RIEMPIE)
                for x, y, tt in ((a, b, t), (b, a, sc.INVERSI.get(t, t))):
                    if tt != rel:
                        continue
                    f = foglio(e["foglio"])
                    k = riga_di(f, e.get("colonna_unita") or idc_us, x)
                    if aggiungi:
                        f.imposta(k, e["colonna"], y)
                    elif _intero(f.valore(k, e["colonna"])) == y:
                        f.imposta(k, e["colonna"], None)
                    else:
                        continue
                    fatto = True
                    break
                if fatto:
                    break
            if not fatto and r.get("modo") == "foglio":
                f = foglio(r["foglio"])
                ca, ct, cb = r["colonne"]
                if aggiungi:
                    f.aggiungi({ca: a, ct: parola(t, [x.get(ct) for x in f.righe.values()]), cb: b})
                    fatto = True
                else:
                    via = [k for k, x in f.righe.items()
                           if _intero(x.get(ca)) is not None and _intero(x.get(cb)) is not None and
                           _diretto(_intero(x.get(ca)), voc.get(_norm(x.get(ct))) or _rapporto(x.get(ct)) or "",
                                    _intero(x.get(cb))) == _diretto(a, t, b)]
                    if via:
                        f.togli(via)
                        fatto = True
            elif not fatto and r.get("modo") == "colonne":
                f = foglio(r["foglio"])
                colonne = {voc.get(_norm(rel)) or _rapporto(rel): col for col, rel in r["colonne"].items()}
                for x, y, tt in ((a, b, t), (b, a, sc.INVERSI.get(t) or sc.DA_INVERSO.get(t))):
                    if tt in colonne:
                        col = colonne[tt]
                        k = riga_di(f, idc_us, x)
                        vecchio = f.valore(k, col)
                        numeri = lista_numeri(vecchio)
                        if aggiungi and y not in numeri:
                            numeri.append(y)
                        elif not aggiungi and y in numeri:
                            numeri.remove(y)
                        else:
                            continue
                        f.imposta(k, col, scrivi_lista(vecchio, numeri))
                        fatto = True
                        break
            elif not fatto and r.get("modo") == "testo":
                f = foglio(r["foglio"])
                cu = r.get("colonna_unita") or idc_us
                for x, y, tt in ((a, b, t), (b, a, sc.INVERSI.get(t) or sc.DA_INVERSO.get(t))):
                    k = f.trova(cu, x, filtri_scheda)
                    if len(k) != 1:
                        continue
                    k = k[0]
                    vecchio = f.valore(k, r["colonna"])
                    nuovo = _riscrivi_testo(vecchio, tt, y, aggiungi, parola)
                    if nuovo is not None:
                        f.imposta(k, r["colonna"], nuovo)
                        fatto = True
                        break
            if not fatto:
                raise ErroreModifica("la ricetta non dice dove scrivere questo rapporto")
            m.update(scritto=True)
            scritte += 1
        except ErroreModifica as e:
            saltate.append((m, str(e)))

    # ---------------------------------------------------------------- salvataggio, con copia
    file_, copie_fatte = [], []
    per_file = {}
    for f in aperti.values():
        if f.cambiato:
            per_file.setdefault(f.path, f)       # un foglio per file basta: salva tutto il file
    if per_file:
        os.makedirs(copie, exist_ok=True)
        stamp = _dt.datetime.now().strftime("%Y%m%d-%H%M%S")
        for path in per_file:
            b, e = os.path.splitext(os.path.basename(path))
            dest = os.path.join(copie, f"{b}.{stamp}{e}")
            shutil.copy2(path, dest)
            copie_fatte.append(dest)
            file_.append(path)
    for path, f in per_file.items():
        if f.ext == ".csv":
            for g in aperti.values():            # un CSV è un foglio solo
                if g.path == path and g.cambiato:
                    g.salva()
        else:
            f.salva()
    for r in risorse.values():
        if hasattr(r, "close") and not hasattr(r, "worksheets"):
            r.close()
    # i file appena scritti non vanno segnalati come «cambiati fuori dall'app»
    for p in file_:
        _aggiorna_impronta(scavo, p)
    scavo.registra("riscrittura", dict(scritte=scritte, saltate=len(saltate), file=[os.path.basename(p) for p in file_]))
    return dict(scritte=scritte, saltate=[(dict(v), motivo) for v, motivo in saltate], file=file_, copie=copie_fatte)


def _riscrivi_testo(vecchio, rel, y, aggiungi, parola):
    """Rapporti scritti come testo: formato di pyArchInit ([['Copre', '2', ...]]) oppure testo libero."""
    from .importa import _rapporti_da_testo, _rapporto, _intero
    s = "" if _vuoto(vecchio) else str(vecchio).strip()
    if s.startswith("[["):
        import ast
        try:
            el = [list(x) for x in ast.literal_eval(s)]
        except (ValueError, SyntaxError):
            return None
        ci = [i for i, x in enumerate(el) if len(x) >= 2 and _rapporto(x[0]) == rel and _intero(x[1]) == y]
        if aggiungi:
            if ci:
                return None
            modello = el[0] if el else [None, None]
            parole = [x[0] for x in el]
            w = parola(rel, parole)
            if w == rel and all(str(p)[:1].isupper() for p in parole):
                w = rel.capitalize()              # pyArchInit scrive «Copre», «Coperto da»
            nuovo = [w, str(y)] + list(modello[2:])
            el.append(nuovo)
        else:
            if not ci:
                return None
            el = [x for i, x in enumerate(el) if i not in ci]
        return str(el)
    coppie = _rapporti_da_testo(s)
    if aggiungi:
        if (rel, y) in coppie:
            return None
        coppie.append((rel, y))
    else:
        if (rel, y) not in coppie:
            return None
        coppie = [c for c in coppie if c != (rel, y)]
    gruppi = {}
    for rr, n in coppie:
        gruppi.setdefault(rr, []).append(n)
    return "; ".join(f"{rr} {', '.join(map(str, ns))}" for rr, ns in gruppi.items()) or None


def _aggiorna_impronta(scavo, path):
    p = os.path.abspath(path)
    for s in scavo.sorgenti:
        if os.path.abspath(s["percorso"]) == p:
            s.update(sha256=_sha256(p), dimensione=os.path.getsize(p))
            return
    scavo.sorgenti.append(dict(percorso=p, tipo="tabella", sha256=_sha256(p), dimensione=os.path.getsize(p),
                               importato=_ora()))


# ---------------------------------------------------------------------------------------------- dai file
def sorgenti_cambiate(scavo):
    """File d'origine cambiati (o spariti) dopo l'importazione o l'ultima riscrittura."""
    out = []
    for s in scavo.sorgenti:
        p = s["percorso"]
        if not os.path.exists(p):
            out.append(dict(percorso=p, stato="mancante"))
        elif os.path.getsize(p) != s.get("dimensione") or _sha256(p) != s.get("sha256"):
            out.append(dict(percorso=p, stato="cambiato"))
    return out


def _impronte(scavo):
    """Per ogni unità, un'impronta dei dati che ne decidono la forma: campi geometrici della scheda,
    pianta, quote, rapporti. (Una descrizione cambiata non chiede di ricostruire nulla.)"""
    imp = {}
    for u, r in {**scavo.schede_us(), **scavo.schede_usm()}.items():
        imp.setdefault(u, []).append(json.dumps({k: _json(r.get(k)) for k in sorted(GEOMETRICI)}, default=str))
    for nome, campo in ((sc.L_US, sc.F_US), (sc.L_USM, sc.F_USM)):
        g = scavo.layers.get(nome)
        if g is not None:
            for u, geom in zip(g[campo], g.geometry):
                if not _vuoto(u):
                    imp.setdefault(int(u), []).append(geom.wkb_hex if geom is not None else "")
    q = scavo.layers.get(sc.L_QUOTE)
    if q is not None and len(q):
        for u, g in q.groupby(sc.F_US):
            pts = sorted((round(p.x, 3), round(p.y, 3), round(p.z, 3) if p.has_z else 0) for p in g.geometry)
            imp.setdefault(int(u), []).append(str(pts))
    rap = scavo.tabelle.get(sc.S_RAPPORTI)
    if rap is not None:
        c = list(rap.columns[:3])
        for a, t, b in zip(rap[c[0]], rap[c[1]], rap[c[2]]):
            if _vuoto(a) or _vuoto(b):
                continue
            x = _diretto(int(a), t, int(b))
            imp.setdefault(x[0], []).append(str(x))
            imp.setdefault(x[2], []).append(str(x))
    return {u: "|".join(sorted(v)) for u, v in imp.items()}


def ricarica(scavo, scarta_modifiche=False):
    """Rilegge i file d'origine con la stessa ricetta. Ritorna (nuovo Scavo, unità da ricostruire, note).

    Parametri, unità escluse, registro e modello restano; si ricostruiscono solo le unità cambiate
    (e quelle che vi si appoggiano sopra)."""
    from . import importa
    if scavo.abbinamento is None:
        raise ErroreModifica("Il progetto non ricorda da quali file viene")
    if in_attesa(scavo) and not scarta_modifiche:
        raise ErroreModifica(f"{len(in_attesa(scavo))} modifiche non ancora scritte nei file d'origine: "
                             "scrivile prima, oppure scartale")
    mancanti = [s["percorso"] for s in sorgenti_cambiate(scavo) if s["stato"] == "mancante"]
    if mancanti:
        raise ErroreModifica("File d'origine non trovati: " + ", ".join(os.path.basename(p) for p in mancanti))
    nuovo = importa.applica(scavo.abbinamento)
    from dataclasses import fields
    for f in fields(scavo.parametri):
        if f.name != "superficie":
            setattr(nuovo.parametri, f.name, getattr(scavo.parametri, f.name))
    prima, dopo = _impronte(scavo), _impronte(nuovo)
    cambiate = {u for u in set(prima) | set(dopo) if prima.get(u) != dopo.get(u)}
    nuovo.meta = dict(scavo.meta, modificato=_ora())
    nuovo.storico = list(scavo.storico)
    nuovo.modifiche = [m for m in scavo.modifiche if m.get("scritto")]
    note = [f"{len(cambiate)} unità cambiate nei file d'origine"] if cambiate else ["Nessuna unità cambiata"]
    ricostruire = set()
    if scavo.modello is not None:
        nuovo.modello = scavo.modello
        vive = set(dopo)
        for u in list(nuovo.modello.unita):
            if u not in vive:
                nuovo.modello.unita.pop(u)
        G = nuovo.rapporti().grafo
        ricostruire = {u for u in cambiate if u in vive}
        for x in list(ricostruire):
            if x in G:
                ricostruire |= nx.ancestors(G, x)
    nuovo.registra("aggiornamento dai file", dict(cambiate=len(cambiate)))
    return nuovo, sorted(ricostruire), note
